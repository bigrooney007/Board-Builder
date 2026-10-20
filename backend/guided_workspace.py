"""Authorized adapters into the existing recruitment/recommitment tools.

Each new paid engagement has its own internal workspace tenant. Its short-lived
delegation token is generated only on the server and never given to a browser.
"""
import asyncio
import io
import json
import os
import re
import secrets
import uuid
from datetime import datetime, timedelta, timezone

import httpx
import jwt
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response
from pydantic import BaseModel, Field

from consultation_routes import no_cache, origin
from guided_services import admin_record, check_service, context, continuation, email_text, new_id, now, owner_id, prepare_recommendations, preview_email, provision_workspace, record_by_token, render, send_message, settings
from workflow_delivery import deliver_once

RECRUITMENT_MATERIALS = ["organization_overview","board_manual","board_member_agreement","confidentiality_agreement","conflict_of_interest_agreement","board_opportunity","board_recruitment_job_post","linkedin_post","social_posts","recruitment_emails","personal_invitation_email","referral_request_email","general_interview_invitation","recruitment_communications","onboarding_script","onboarding_agenda"]
APPLICANT_RESOURCES = {'interview_guide','interview_invitation','before_interview_rejection','after_interview_thank_you','candidate_referee_request','reference_request_email','referee_confirmation_email','conditional_offer','unconditional_offer','formal_appointment_letter','onboarding_email','board_member_engagement_guide','board_member_portfolio','board_member_agreement','confidentiality_agreement','conflict_of_interest_agreement'}

async def authorize_delegation(db,request,payload,member):
    service=member.get("guided_service_workspace")
    if not service:return member
    prefix="/api/workspace/" if service=="recruitment" else "/api/game/" if service=="board-fundraising" else "/api/reactivation/"
    if payload.get("guided_service")!=service or not request.url.path.startswith(prefix):raise HTTPException(403,"This service access is limited to its own workspace tools.")
    row=await db.guided_service_customers.find_one({"customer_id":member.get("guided_customer_id"),"service":service,"workspace_user_id":member["user_id"],"payment_status":"paid"},{"_id":0,"customer_id":1})
    if not row:raise HTTPException(403,"A verified paid engagement is required for this workspace.")
    return {**member,"email":member["notification_email"],"guided_facilitator":True,"fundraising_facilitator":service=="board-fundraising"}

async def delegated_request(app,db,row,method,path,payload=None,query=""):
    await provision_workspace(db,row)
    member=await db.members.find_one({"user_id":owner_id(row)},{"_id":0})
    token=jwt.encode({"sub":member["user_id"],"email":member["email"],"role":"member","type":"access","guided_service":row["service"],"exp":datetime.now(timezone.utc)+timedelta(minutes=3)},os.environ["JWT_SECRET"],algorithm="HS256")
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url="http://guided-internal",timeout=180) as client:
        response=await client.request(method,"/api/"+path+("?"+query if query else ""),json=payload,headers={"Authorization":"Bearer "+token})
    if response.status_code>=400:
        try:detail=response.json().get("detail","This workspace action could not be completed.")
        except Exception:detail="This workspace action could not be completed."
        raise HTTPException(response.status_code,detail)
    return response

def current_version(material):
    return next((v for v in material.get("versions",[]) if v["version"]==material.get("current_version")),{})

async def sync_profiles(db,row):
    if row.get("payment_status")!="paid" or not row.get("profiles_approved"):return
    from workspace_service import save_generation,get_current_material
    old=await get_current_material(db,owner_id(row),"powerhouse_board_blueprint")
    roles=row.get("recommendations") or []
    if not old or ((old.get("current") or {}).get("structured") or {}).get("priority_roles")!=roles:
        material=await save_generation(db,owner_id(row),"powerhouse_board_blueprint",{"priority_roles":roles},"Board member profiles reviewed and agreed during the guided recruitment process.")
        await db.generated_materials.update_one({"material_id":material["material_id"]},{"$set":{"status":"Approved","approved_at":now()}})
    await db.funnel_leads.update_one({"lead_id":row["customer_id"]},{"$set":{"answers.new_members_needed":str(row.get("recruitment_count") or len(roles))}})

async def work_tick(db,app,only_id=None):
    from guided_service_defaults import ACTIVE_SERVICES
    query={"service":{"$in":list(ACTIVE_SERVICES)},"submitted_at":{"$exists":True}}
    if only_id:query["customer_id"]=only_id
    rows=await db.guided_service_customers.find(query,{"_id":0}).sort("work_checked_at",1).limit(30).to_list(30)
    for row in rows:
        await db.guided_service_customers.update_one({"customer_id":row["customer_id"]},{"$set":{"work_checked_at":now()}})
        if row["service"]=="recruitment" and (row.get("analysis") or {}).get("status") in {"queued","failed","generating"} and row.get("analysis_attempts",0)<3:await prepare_recommendations(db,row)
        job=row.get("material_job") or {}
        if row.get("payment_status")!="paid" or job.get("status") not in {"queued","generating"}:continue
        claim=await db.guided_service_customers.update_one({"customer_id":row["customer_id"],"$or":[{"material_job.lock_until":{"$exists":False}},{"material_job.lock_until":{"$lt":now()}}]},{"$set":{"material_job.status":"generating","material_job.lock_until":(datetime.now(timezone.utc)+timedelta(minutes=20)).isoformat()}})
        if not claim.modified_count:continue
        failures=[]
        for kind in job.get("types",[]):
            existing=await db.generated_materials.find_one({"user_id":owner_id(row),"type":kind,"application_id":""},{"_id":0})
            if existing and existing.get("status") in {"Approved","Edited","Ready","Draft","Generated","Needs Review"}:continue
            try:
                await delegated_request(app,db,row,"POST","workspace/generate",{"type":kind,"instructions":"Use the human-approved board member profiles and the organization's actual information. Do not invent financial figures, dates, commitments, contacts or background-check results. Preserve all approved profiles."})
            except Exception as exc:
                failures.append({"type":kind,"error":str(getattr(exc,"detail","Generation could not finish. Try this document again."))[:300]})
        await db.guided_service_customers.update_one({"customer_id":row["customer_id"]},{"$set":{"material_job.status":"needs_attention" if failures else "complete","material_job.errors":failures,"material_job.completed_at":now()},"$unset":{"material_job.lock_until":""}})

async def work_loop(db,app):
    import logging
    while True:
        try:await work_tick(db,app)
        except Exception:logging.getLogger(__name__).exception("Guided service work will retry")
        await asyncio.sleep(10)

class Profiles(BaseModel):
    roles:list[dict]=Field(min_length=1,max_length=50)
    count:int=Field(ge=1,le=50)
    approved:bool=False

def allowed_tool(service,method,path,admin):
    if not re.fullmatch(r"[A-Za-z0-9_/-]{1,250}",path) or ".." in path:return False
    if service=="recruitment":
        if not path.startswith("workspace/"):return False
        short=path.removeprefix("workspace/")
        if method=="GET":return bool(re.fullmatch(r"(?:profile|board-profiles|materials(?:/[\w-]+)?|opportunity|applications(?:/[\w-]+(?:/cv|/onboarding-conclusion)?)?|interview-scheduling|onboarding-session|board-profile-form|signatures(?:/[\w-]+/download)?)",short))
        if not admin:return bool(method=="POST" and re.fullmatch(r"(?:applications/[\w-]+/interview-guide/retry|generate|materials/[\w-]+/approve|signatures/prepare)",short) or method=="PATCH" and re.fullmatch(r"applications/[\w-]+",short) or method=="PUT" and re.fullmatch(r"(?:applications/[\w-]+/onboarding-conclusion|materials/[\w-]+|onboarding-session)",short))
        return bool(method in {"PUT","POST","PATCH"} and re.fullmatch(r"(?:profile|profile/confirm|strategy-intake|board-profiles|generate|materials/[\w-]+(?:/approve|/share)?|interview-scheduling|onboarding-session|opportunity/application(?:/generate)?|opportunity/close|opportunity/broadcast/retry|applications/[\w-]+(?:/references|/role-recommendation|/interview-guide/retry|/onboarding-conclusion)?|signatures/prepare|signatures/[\w-]+/send)",short))
    if service=="board-recommitment":
        if not path.startswith("reactivation/"):return False
        short=path.removeprefix("reactivation/")
        if method=="GET":return bool(re.fullmatch(r"(?:setup|branding|roster|step3|my-board|understand|board-summary|board-members/[\w-]+/(?:response(?:/pdf)?|call-script|portfolio-email|outcome-email)|materials/[\w-]+(?:/pdf)?|recommitment-form|recommitment-email-draft)",short))
        if not admin:return False
        return bool(method in {"PUT","POST"} and re.fullmatch(r"(?:setup|branding|board-members|board-members/import|board-members/[\w-]+/(?:call-notes|conversation-script|conclusion|outcome|confirmed-role|portfolio|analysis|direction|outcome-email)|materials/[\w-]+(?:/approve|/approve-portfolio)?|recommitment-form(?:/generate|/approve)?|recommitment-email-draft(?:/approve)?)",short))
    return False

async def material_snapshot(db,row,material_id,approved=False):
    material=await db.generated_materials.find_one({"material_id":material_id,"user_id":owner_id(row)},{"_id":0})
    if not material:raise HTTPException(404,"Resource not found for this organization.")
    if approved and material.get("status")!="Approved":raise HTTPException(403,"This resource has not been approved for release.")
    return material,current_version(material)

async def notify_guided_owner(db,user_id,event,subject,body,query="",template_key=None,variables=None):
    row=await db.guided_service_customers.find_one({"workspace_user_id":user_id,"payment_status":"paid"},{"_id":0})
    if not row:return None
    cta="VIEW YOUR RESOURCES"
    if not template_key:
        prefixes={"guided-signed-":"signature_completed","guided-member-profile-completed":"profile_response_notice","guided-interview-guide":"interview_guide_ready","guided-campaign-launched":"campaign_launched","guided-approved-conversation":"conversation_guide_approved"}
        template_key=next((value for prefix,value in prefixes.items() if event.startswith(prefix)),None)
    if template_key:
        config,_,_=await settings(db,row["service"]);message=render(config["operational_emails"][template_key],{**context(row),**(variables or {})});subject=message["subject"];body=email_text(message);cta=message["cta"]
    return await deliver_once(db,user_id,event,row["email"],subject,body,continuation(row,True)+query,cta_label=cta)

def create_guided_workspace_router(db):
    router=APIRouter(prefix="/api",dependencies=[Depends(no_cache)])

    @router.api_route("/admin/services/{service}/customers/{ident}/tool/{path:path}",methods=["GET","PUT","POST","PATCH"])
    async def admin_tool(service:str,ident:str,path:str,request:Request):
        row=await admin_record(db,service,ident,request,paid=True)
        if not allowed_tool(service,request.method,path,True):raise HTTPException(403,"This action is outside the selected service.")
        if path=="workspace/generate":
            payload=await request.json()
            if payload.get("type") not in RECRUITMENT_MATERIALS+['interview_guide','interview_invitation','before_interview_rejection','conditional_offer','unconditional_offer','formal_appointment_letter','onboarding_email','board_member_portfolio','ninety_day_plan','after_interview_thank_you','candidate_referee_request','referee_confirmation_email','reference_call_script','reference_evaluation_form','reference_request_email','board_member_engagement_guide']:raise HTTPException(422,"Choose a recruitment resource supported by this service.")
        else:payload=await request.json() if request.method!="GET" else None
        if path=="workspace/signatures/prepare":
            existing=await reviewed_signature(row,payload)
            if existing:return existing
        if request.method=="POST" and re.fullmatch(r"workspace/signatures/[\w-]+/send",path):raise HTTPException(403,"Use the reviewed signature delivery action.")
        response=await delegated_request(request.app,db,row,request.method,path,payload,request.url.query)
        return Response(response.content,response.status_code,headers={"Cache-Control":"no-store",**({"Content-Disposition":response.headers['content-disposition']} if 'content-disposition' in response.headers else {})},media_type=response.headers.get("content-type","application/json"))

    async def reviewed_signature(row,payload):
        from workspace_service import get_current_material
        material=await get_current_material(db,owner_id(row),payload.get("agreement_type"),payload.get("application_id")) or await get_current_material(db,owner_id(row),payload.get("agreement_type"),"")
        if not material or material["material"].get("status")!="Approved":raise HTTPException(409,"Approve the exact agreement before preparing a signature request.")
        return await db.signature_requests.find_one({"owner_user_id":owner_id(row),"application_id":payload.get("application_id"),"material_id":material["material"]["material_id"],"agreement_version":material["current"]["version"]},{"_id":0,"token":0})

    @router.api_route("/services/{service}/customer/{token}/tool/{path:path}",methods=["GET","PUT","POST","PATCH"])
    async def customer_tool(service:str,token:str,path:str,request:Request):
        row=await record_by_token(db,service,token,paid=True)
        if not allowed_tool(service,request.method,path,False):raise HTTPException(403,"This action is not available from your customer resources.")
        payload=await request.json() if request.method!="GET" else None
        if "materials" in path:
            match=re.fullmatch(r"workspace/materials/([\w-]+)(?:/approve)?",path)
            if not match:raise HTTPException(403,"Use your approved resource list.")
            material,_=await material_snapshot(db,row,match[1],approved=request.method=="GET")
            if request.method!="GET" and not material.get("application_id"):raise HTTPException(403,"Approved campaign resources remain unchanged. These editing tools are for your applicant-specific documents.")
        if path=="workspace/generate":
            if payload.get("type") not in APPLICANT_RESOURCES or not payload.get("application_id"):raise HTTPException(403,"Choose an applicant-specific interview, reference, appointment or onboarding resource.")
            if not await db.opportunity_applications.find_one({"application_id":payload["application_id"],"owner_user_id":owner_id(row)},{"_id":1}):raise HTTPException(404,"Applicant not found for this organization.")
        if path=="workspace/signatures/prepare":
            existing=await reviewed_signature(row,payload)
            if existing:return existing
        response=await delegated_request(request.app,db,row,request.method,path,payload,request.url.query)
        return Response(response.content,response.status_code,media_type=response.headers.get("content-type","application/json"),headers={"Cache-Control":"no-store",**({"Content-Disposition":response.headers['content-disposition']} if 'content-disposition' in response.headers else {})})

    async def recruitment_actor(request,ident=None,token=None):
        return await admin_record(db,"recruitment",ident,request,paid=True) if ident else await record_by_token(db,"recruitment",token,paid=True)

    async def candidate_message(row,application_id,action):
        application=await db.opportunity_applications.find_one({"application_id":application_id,"owner_user_id":owner_id(row)},{"_id":0})
        if not application:raise HTTPException(404,"Applicant not found for this organization.")
        if action not in {"invite","reject"}:raise HTTPException(422,"Choose an interview invitation or rejection.")
        details=row.get("interview_details") or {}
        if action=="invite" and not (details.get("scheduling_link") or details.get("availability")):raise HTTPException(409,"Save the organization's confirmed interview arrangements first.")
        config,_,_=await settings(db,"recruitment")
        values={**context(row),"applicant_name":(application.get("profile_snapshot") or {}).get("full_name") or application["applicant_email"],"interview_details":"\n".join(str(details.get(k,"")) for k in ("scheduling_link","availability","preferred_method","contact_person","instructions") if details.get(k))}
        return application,render(config["operational_emails"]["interview_invite" if action=="invite" else "applicant_reject"],values)

    @router.post("/admin/services/recruitment/customers/{ident}/applicants/{application_id}/profile-form")
    @router.post("/services/recruitment/customer/{token}/applicants/{application_id}/profile-form")
    async def applicant_profile(application_id:str,payload:dict,request:Request,ident:str=None,token:str=None):
        row=await recruitment_actor(request,ident,token)
        application=await db.opportunity_applications.find_one({"application_id":application_id,"owner_user_id":owner_id(row)},{"_id":0})
        if not application:raise HTTPException(404,"Applicant not found for this organization.")
        link=await db.board_profile_links.find_one({"user_id":owner_id(row),"application_id":application_id},{"_id":0})
        if not link:
            snapshot=application.get("profile_snapshot") or {}
            link={"user_id":owner_id(row),"application_id":application_id,"token":secrets.token_urlsafe(32),"prefill":{"full_name":snapshot.get("full_name",""),"email":application["applicant_email"],"professional_title":snapshot.get("profession","")},"status":"Created","created_at":now()}
            await db.board_profile_links.insert_one(link.copy())
        url=origin()+"/board-profile/"+link["token"]
        if payload.get("approved"):
            config,_,_=await settings(db,row["service"]);message=render(config["operational_emails"]["board_profile_invitation"],context(row))
            return {"url":url,**await deliver_once(db,owner_id(row),"guided-member-profile-"+application_id,application["applicant_email"],message["subject"],email_text(message),url,cta_label=message["cta"])}
        return {"url":url,"recipient":application["applicant_email"]}

    @router.post("/admin/services/recruitment/customers/{ident}/applicants/{application_id}/preview")
    @router.post("/services/recruitment/customer/{token}/applicants/{application_id}/preview")
    async def applicant_preview(application_id:str,payload:dict,request:Request,ident:str=None,token:str=None):
        row=await recruitment_actor(request,ident,token);application,message=await candidate_message(row,application_id,payload.get("action"))
        opportunity=await db.opportunities.find_one({"user_id":owner_id(row)},{"_id":0}) or {}
        url=origin()+"/board-opportunities/"+opportunity.get("slug","")+"/apply"
        return {**message,"recipient":application["applicant_email"],"action":payload["action"],"application_status":application["status"],"html":preview_email(message,url)}

    @router.post("/admin/services/recruitment/customers/{ident}/applicants/{application_id}/send")
    @router.post("/services/recruitment/customer/{token}/applicants/{application_id}/send")
    async def applicant_send(application_id:str,payload:dict,request:Request,ident:str=None,token:str=None):
        row=await recruitment_actor(request,ident,token);application,default=await candidate_message(row,application_id,payload.get("action"))
        if not payload.get("approved"):raise HTTPException(409,"Review and approve this communication before sending.")
        request_id=str(uuid.UUID(payload["request_id"]))
        previous=await db.workflow_email_deliveries.find_one({"user_id":owner_id(row),"event_key":"guided-candidate-"+application_id+"-"+request_id,"recipient":application["applicant_email"]},{"_id":0})
        if previous and previous.get("status") in {"sent","captured"}:return {"status":previous["status"],"recipient":application["applicant_email"]}
        if application["status"]!=payload.get("application_status"):raise HTTPException(409,"The applicant's status changed. Review the invitation again.")
        message={k:str(payload.get(k,default[k]))[:20000] for k in default}
        request_id=str(uuid.UUID(payload["request_id"]))
        opportunity=await db.opportunities.find_one({"user_id":owner_id(row)},{"_id":0}) or {}
        result=await deliver_once(db,owner_id(row),"guided-candidate-"+application_id+"-"+request_id,application["applicant_email"],message["subject"],email_text(message),origin()+"/board-opportunities/"+opportunity.get("slug","")+"/apply",cta_label=message["cta"])
        if result["status"] in {"sent","captured"}:
            status="Interview Invited" if payload["action"]=="invite" else "Not Moving to Interview"
            # Use the exact stored status identifiers from the original application system.
            from workspace_service import APPLICATION_STATUSES
            status=next((v for v in APPLICATION_STATUSES if v.replace(" ","").lower()==status.replace(" ","").lower()),status)
            await db.opportunity_applications.update_one({"application_id":application_id,"owner_user_id":owner_id(row)},{"$set":{"status":status,"updated_at":now()}})
            if payload["action"]=="invite":
                guide=await delegated_request(request.app,db,row,"POST","workspace/applications/"+application_id+"/interview-guide/retry",{})
                if (guide.json().get("interview_guide") or {}).get("status")=="Ready":await notify_guided_owner(db,owner_id(row),"guided-interview-guide-"+application_id,"Your Applicant Interview Guide","An interview guide is available with the application. Review it before your conversation.","?applicant="+application_id)
        return result

    @router.post("/admin/services/recruitment/customers/{ident}/signatures/{signature_id}/send")
    @router.post("/services/recruitment/customer/{token}/signatures/{signature_id}/send")
    async def send_signature(signature_id:str,payload:dict,request:Request,ident:str=None,token:str=None):
        row=await recruitment_actor(request,ident,token)
        item=await db.signature_requests.find_one({"request_id":signature_id,"owner_user_id":owner_id(row)},{"_id":0})
        if not item:raise HTTPException(404,"Agreement not found for this organization.")
        if not payload.get("approved"):raise HTTPException(409,"Review and approve sending this agreement.")
        if item.get("status")=="Signed":return {"status":"already_signed"}
        config,_,_=await settings(db,row["service"]);message=render(config["operational_emails"]["signature_invitation"],context(row))
        result=await deliver_once(db,owner_id(row),"guided-signature-invitation-"+signature_id,item["board_member_email"],message["subject"],email_text(message),origin()+"/sign/"+item["token"],cta_label=message["cta"])
        if result["status"] in {"sent","captured"}:await db.signature_requests.update_one({"request_id":signature_id,"status":{"$ne":"Signed"}},{"$set":{"status":"Sent","sent_at":now()}})
        return result

    @router.put("/admin/services/{service}/customers/{ident}/organization")
    async def organization(service:str,ident:str,payload:dict,request:Request):
        row=await admin_record(db,service,ident,request,paid=True)
        fields={key:str(payload.get(key,""))[:12000] for key in ("organization_name","mission","website","founder_title","present_board","priorities")}
        if not fields["organization_name"].strip() or not fields["mission"].strip():raise HTTPException(422,"Confirm the organization's name and mission.")
        await db.recruitment_profiles.update_one({"user_id":owner_id(row)},{"$set":{"data."+k:v for k,v in fields.items()}})
        await db.recruitment_profiles.update_one({"user_id":owner_id(row)},{"$set":{"confirmed":True,"updated_at":now()}})
        await db.guided_service_customers.update_one({"customer_id":ident},{"$set":{"organization":fields["organization_name"],"clarified_details":fields,"updated_at":now()}})
        return {"status":"saved"}

    @router.put("/admin/services/recruitment/customers/{ident}/profiles")
    async def profiles(ident:str,payload:Profiles,request:Request):
        row=await admin_record(db,"recruitment",ident,request)
        if await db.opportunities.find_one({"user_id":owner_id(row),"status":{"$in":["Published","Closed"]}}):raise HTTPException(409,"The launched campaign's agreed profiles are preserved.")
        if len(payload.roles)!=payload.count or any(not str(r.get("role_name","")).strip() for r in payload.roles):raise HTTPException(422,"Give each profile a title and match the agreed number of board members.")
        roles=[]
        for raw in payload.roles:
            role={k:str(raw.get(k,""))[:6000] for k in ("role_name","why_this_person_is_important","how_this_person_can_support")}
            for key in ("what_to_look_for","gap_this_role_fills"):
                values=raw.get(key) or []
                if not isinstance(values,(str,list)):raise HTTPException(422,"Use text or a list of criteria for each profile.")
                role[key]=[str(v)[:1200] for v in (values.splitlines() if isinstance(values,str) else values)][:30]
            roles.append(role)
        if row.get("recommendations")==roles and row.get("recruitment_count")==payload.count and row.get("profiles_approved")==payload.approved:return {"status":"approved" if payload.approved else "saved"}
        updates={"recommendations":roles,"recruitment_count":payload.count,"profiles_approved":payload.approved,"profiles_updated_at":now()}
        await db.guided_service_customers.update_one({"customer_id":ident},{"$set":updates,"$push":{"profile_history":{"roles":roles,"approved":payload.approved,"created_at":now()}}})
        await db.generated_materials.update_many({"user_id":owner_id(row),"type":{"$in":RECRUITMENT_MATERIALS}},{"$set":{"status":"Needs Review"}})
        if row.get("payment_status")=="paid":await sync_profiles(db,{**row,**updates})
        return {"status":"approved" if payload.approved else "saved"}

    @router.post("/admin/services/recruitment/customers/{ident}/generate-materials")
    async def generate_materials(ident:str,request:Request):
        row=await admin_record(db,"recruitment",ident,request,paid=True)
        if not row.get("profiles_approved"):raise HTTPException(409,"Approve the agreed board member profiles first.")
        await sync_profiles(db,row)
        # The original generator creates an organization-specific application as well.
        await delegated_request(request.app,db,row,"POST","workspace/opportunity/application/generate",{})
        await db.guided_service_customers.update_one({"customer_id":ident},{"$set":{"material_job":{"status":"queued","types":RECRUITMENT_MATERIALS,"queued_at":now()}}})
        return {"status":"queued","types":RECRUITMENT_MATERIALS}

    @router.put("/admin/services/recruitment/customers/{ident}/interview-details")
    async def interview_details(ident:str,payload:dict,request:Request):
        row=await admin_record(db,"recruitment",ident,request,paid=True)
        if payload.get("method") not in {"link","availability"}:raise HTTPException(422,"Choose a booking link or confirmed availability.")
        value={key:str(payload.get(key,""))[:5000] for key in ("method","scheduling_link","availability","preferred_method","contact_person","instructions")}
        await delegated_request(request.app,db,row,"PUT","workspace/interview-scheduling",value)
        await db.guided_service_customers.update_one({"customer_id":ident},{"$set":{"interview_details":value}})
        return {"status":"saved"}

    @router.post("/admin/services/recruitment/customers/{ident}/launch")
    async def launch(ident:str,payload:dict,request:Request):
        row=await admin_record(db,"recruitment",ident,request,paid=True)
        if not payload.get("confirmed"):raise HTTPException(409,"Review the final campaign preview and confirm publication.")
        existing=await db.opportunities.find_one({"user_id":owner_id(row)},{"_id":0})
        if existing and existing.get("status")=="Published":return {"status":"Published","opportunity":existing,"broadcast_status":existing.get("broadcast_status","Not requested")}
        response=await delegated_request(request.app,db,row,"POST","workspace/opportunity/publish",{"network_consent":bool(payload.get("network_consent"))})
        await db.guided_service_customers.update_one({"customer_id":ident},{"$set":{"campaign_launched_at":now()}})
        await notify_guided_owner(db,owner_id(row),"guided-campaign-launched","Your Board Recruitment Campaign Is Published","Your approved board opportunity is published on Nonprofit Board Builder and its application form is ready. Open your resources to copy your campaign link and promotional materials. Your network announcement status is available with the campaign. Publication on other websites requires sharing through your own channels.")
        return response.json()

    @router.get("/services/{service}/customer/{token}/signed/{signature_id}")
    async def signed_download(service:str,token:str,signature_id:str):
        row=await record_by_token(db,service,token,paid=True)
        item=await db.signature_requests.find_one({"owner_user_id":owner_id(row),"request_id":signature_id,"status":"Signed"},{"_id":0})
        if not item:raise HTTPException(404,"Signed agreement not found for this organization.")
        signed=item.get("signed") or {}
        text=item["document_snapshot"]+"\n\nSIGNED ELECTRONICALLY\n"+"\n".join(str(signed.get(k,"")) for k in ("typed_signature","email","date","signed_at"))
        return Response(text.encode(),media_type="text/plain",headers={"Content-Disposition":'attachment; filename="Signed-Board-Agreement.txt"',"Cache-Control":"no-store"})

    @router.get("/services/{service}/customer/{token}/resources")
    async def resources(service:str,token:str):
        row=await record_by_token(db,service,token,paid=True)
        documents=await db.generated_materials.find({"user_id":owner_id(row),"status":"Approved"},{"_id":0}).to_list(500)
        # Only the leader's paid capability can retrieve this organization-wide collection.
        materials=[{"material_id":m["material_id"],"title":m.get("title",m["type"]),"type":m["type"],"application_id":m.get("application_id",""),"text":current_version(m).get("display_text","")} for m in documents]
        opportunity=await db.opportunities.find_one({"user_id":owner_id(row),"status":{"$in":["Published","Closed"]}},{"_id":0})
        signatures=await db.signature_requests.find({"owner_user_id":owner_id(row),"status":"Signed"},{"_id":0,"token":0}).to_list(100)
        return {"organization":row["organization"],"branding":row.get("branding",{}),"materials":materials,"campaign":opportunity,"signatures":signatures}

    async def download(row,material_id,format):
        material,version=await material_snapshot(db,row,material_id,approved=True)
        text=version.get("display_text","");title=material.get("title", "Resource");filename=re.sub(r"[^a-zA-Z0-9-]","-",title)[:80]
        branding=row.get("branding",{})
        if format=="pdf":
            from document_renderer import render_document_pdf
            data=render_document_pdf(title=title,org_name=row["organization"],body_text=text,branding=branding);mime="application/pdf"
        elif format=="docx":
            import base64
            from docx import Document
            from docx.shared import Inches,RGBColor
            doc=Document()
            if branding.get("logo_data"):
                from PIL import Image
                logo=Image.open(io.BytesIO(base64.b64decode(branding["logo_data"].split(",")[1])));buffered=io.BytesIO();logo.convert("RGBA").save(buffered,format="PNG");buffered.seek(0)
                doc.add_picture(buffered,width=Inches(1))
            heading=doc.add_heading(title,0)
            for run in heading.runs:run.font.color.rgb=RGBColor.from_string(branding.get("primary_color","#334155").lstrip('#'))
            doc.add_paragraph(row["organization"])
            for paragraph in text.split('\n'):doc.add_paragraph(paragraph)
            buffer=io.BytesIO();doc.save(buffer);data=buffer.getvalue();mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif format=="txt":data=text.encode();mime="text/plain"
        else:raise HTTPException(422,"Choose PDF, editable Word or text.")
        return Response(data,media_type=mime,headers={"Content-Disposition":f'attachment; filename="{filename}.{format}"',"Cache-Control":"no-store"})

    @router.get("/services/{service}/customer/{token}/materials/{material_id}/download")
    async def customer_download(service:str,token:str,material_id:str,format:str="pdf"):return await download(await record_by_token(db,service,token,paid=True),material_id,format)

    @router.get("/admin/services/{service}/customers/{ident}/materials/{material_id}/download")
    async def admin_download(service:str,ident:str,material_id:str,request:Request,format:str="pdf"):return await download(await admin_record(db,service,ident,request,paid=True),material_id,format)

    @router.post("/admin/services/{service}/customers/{ident}/delivery-preview")
    async def delivery_preview(service:str,ident:str,request:Request):
        row=await admin_record(db,service,ident,request,paid=True);config,_,_=await settings(db,service)
        message=render(config["delivery_email"],context(row));return {**message,"expected_payment_status":row["payment_status"],"html":preview_email(message,continuation(row,True))}

    @router.post("/admin/services/{service}/customers/{ident}/deliver")
    async def deliver(service:str,ident:str,payload:dict,request:Request):
        row=await admin_record(db,service,ident,request,paid=True)
        if not payload.get("approved"):raise HTTPException(409,"Review and approve the resource delivery email.")
        if service=="recruitment" and not await db.opportunities.find_one({"user_id":owner_id(row),"status":"Published"}):raise HTTPException(409,"Launch the recruitment campaign before sending the final resource package.")
        config,_,_=await settings(db,service);message={k:str(payload.get(k,render(config["delivery_email"],context(row))[k])) for k in ("subject","headline","body","cta","support","signature")}
        revision=await db.generated_materials.count_documents({"user_id":owner_id(row),"status":"Approved"})
        result=await send_message(db,row,f"guided-resource-delivery-{payload.get('request_id') or revision}",message,continuation(row,True))
        if result["status"] in {"sent","captured"}:await db.guided_service_customers.update_one({"customer_id":ident},{"$set":{"resource_delivery_at":now()}})
        return result

    from service_handover import register_operator_routes
    register_operator_routes(router,"/admin/services/{service}/customers/{ident}","/services/{service}/operator/{ident}")
    register_operator_routes(router,"/admin/services/recruitment/customers/{ident}","/services/recruitment/operator/{ident}")
    return router
