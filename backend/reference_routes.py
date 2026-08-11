"""Founder-only Recruitment Execution Reference Library (admin access only)."""
import secrets
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from pydantic import BaseModel

from ai_service import extract_cv_text
from auth_service import authenticate_admin


def extract_reference_text(content: bytes, filename: str) -> str:
    import io
    import os as _os
    extension = _os.path.splitext(filename or "")[1].lower()
    try:
        if extension == ".docx":
            import docx
            document = docx.Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)[:120000]
        if extension == ".txt":
            return content.decode("utf-8", errors="ignore")[:120000]
    except Exception:
        pass
    return extract_cv_text(content, filename)

MODULE_LABELS = {0: "General / all modules", 1: "Identify the Board Members Your Organization Needs", 2: "Build Your Recruitment Strategy", 3: "Launch Your Recruitment Campaign", 4: "Interview Your Applicants", 5: "Complete References and Background Checks", 6: "Onboard Your New Board Members"}


class ReferenceUpdate(BaseModel):
    approved: Optional[bool] = None
    module: Optional[int] = None
    resource_types: Optional[list] = None
    title: Optional[str] = None


def create_reference_router(db) -> APIRouter:
    router = APIRouter(prefix="/api/admin/reference-library")

    @router.post("", status_code=201)
    async def upload_reference(request: Request, file: UploadFile = File(...), title: str = Form(""), module: int = Form(0), resource_types: str = Form("")):
        await authenticate_admin(request, db)
        if module not in MODULE_LABELS:
            raise HTTPException(status_code=422, detail="Module must be 0 (general) or 1-6")
        content = await file.read()
        if len(content) > 15 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="File is too large (15MB maximum)")
        text = extract_reference_text(content, file.filename)
        if not text.strip():
            raise HTTPException(status_code=422, detail="Could not extract text from this file. Supported: .docx, .pdf, .doc, .txt")
        record = {
            "reference_id": secrets.token_hex(8), "title": (title or file.filename).strip(),
            "filename": file.filename, "module": module,
            "resource_types": [tag.strip() for tag in resource_types.split(",") if tag.strip()],
            "content_text": text, "approved": False,
            "uploaded_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.reference_materials.insert_one(record.copy())
        record.pop("content_text")
        return {"reference": record, "extracted_characters": len(text)}

    @router.get("")
    async def list_references(request: Request):
        await authenticate_admin(request, db)
        docs = await db.reference_materials.find({}, {"_id": 0, "content_text": 0}).sort("uploaded_at", -1).to_list(200)
        return {"references": docs, "module_labels": MODULE_LABELS}

    @router.get("/{reference_id}")
    async def get_reference(reference_id: str, request: Request):
        await authenticate_admin(request, db)
        doc = await db.reference_materials.find_one({"reference_id": reference_id}, {"_id": 0})
        if not doc:
            raise HTTPException(status_code=404, detail="Reference material not found")
        return {"reference": doc}

    @router.patch("/{reference_id}")
    async def update_reference(reference_id: str, payload: ReferenceUpdate, request: Request):
        await authenticate_admin(request, db)
        update = {k: v for k, v in payload.model_dump().items() if v is not None}
        if "module" in update and update["module"] not in MODULE_LABELS:
            raise HTTPException(status_code=422, detail="Module must be 0 (general) or 1-6")
        if not update:
            raise HTTPException(status_code=422, detail="Nothing to update")
        result = await db.reference_materials.update_one({"reference_id": reference_id}, {"$set": update})
        if not result.matched_count:
            raise HTTPException(status_code=404, detail="Reference material not found")
        doc = await db.reference_materials.find_one({"reference_id": reference_id}, {"_id": 0, "content_text": 0})
        return {"reference": doc}

    @router.delete("/{reference_id}")
    async def delete_reference(reference_id: str, request: Request):
        await authenticate_admin(request, db)
        result = await db.reference_materials.delete_one({"reference_id": reference_id})
        if not result.deleted_count:
            raise HTTPException(status_code=404, detail="Reference material not found")
        return {"status": "deleted"}

    return router
